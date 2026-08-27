import calendar
from datetime import datetime, date, time, timedelta
from io import BytesIO
from typing import List, Optional, Dict, Any

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, and_, or_, desc
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.core.database import get_db
from app.models.models import (
    Payment, Expense, User, UserRole, Group, GroupStudent, Course,
    TeacherPayroll, StudentBilling, Attendance, AttendanceStatus,
    Lesson, Exam, ExamResult, Lead, LeadStatus
)
from app.api.deps import get_current_admin

router = APIRouter()

MONTH_NAMES_UZ = {
    1: "Yanvar", 2: "Fevral", 3: "Mart", 4: "Aprel",
    5: "May", 6: "Iyun", 7: "Iyul", 8: "Avgust",
    9: "Sentyabr", 10: "Oktyabr", 11: "Noyabr", 12: "Dekabr"
}

def resolve_date_range(
    period_type: str = "monthly",
    year: Optional[int] = None,
    month: Optional[int] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None
):
    today = date.today()
    if period_type == "yearly":
        selected_year = year or today.year
        start_d = date(selected_year, 1, 1)
        end_d = date(selected_year, 12, 31)
        label = f"{selected_year}-yil (Yillik hisobot)"
    elif period_type == "custom":
        start_d = start_date or date(today.year, today.month, 1)
        end_d = end_date or today
        if start_d > end_d:
            start_d, end_d = end_d, start_d
        label = f"{start_d.strftime('%d.%m.%Y')} — {end_d.strftime('%d.%m.%Y')}"
    else:  # monthly
        selected_year = year or today.year
        selected_month = month or today.month
        last_day = calendar.monthrange(selected_year, selected_month)[1]
        start_d = date(selected_year, selected_month, 1)
        end_d = date(selected_year, selected_month, last_day)
        month_name = MONTH_NAMES_UZ.get(selected_month, f"{selected_month}-oy")
        label = f"{selected_year}-yil {month_name} oyi"

    start_dt = datetime.combine(start_d, time.min)
    end_dt = datetime.combine(end_d, time.max)
    return start_d, end_d, start_dt, end_dt, label


async def collect_report_data(
    db: AsyncSession,
    start_d: date,
    end_d: date,
    start_dt: datetime,
    end_dt: datetime,
    period_type: str,
    label: str
) -> Dict[str, Any]:
    # 1. Payments in date range
    pay_stmt = (
        select(Payment, User.full_name, User.login_id, User.phone)
        .outerjoin(User, Payment.student_id == User.id)
        .where(
            func.date(Payment.created_at) >= start_d,
            func.date(Payment.created_at) <= end_d
        )
        .order_by(desc(Payment.created_at))
    )
    pay_res = await db.execute(pay_stmt)
    payment_rows = pay_res.all()

    total_income = 0.0
    by_payment_method: Dict[str, float] = {
        "CASH": 0.0, "CARD": 0.0, "CLICK": 0.0, "PAYME": 0.0, "UZUM": 0.0, "BANK_TRANSFER": 0.0
    }
    payments_list = []
    for p, s_name, s_login, s_phone in payment_rows:
        total_income += p.amount
        method = p.payment_method.value if hasattr(p.payment_method, "value") else str(p.payment_method)
        by_payment_method[method] = by_payment_method.get(method, 0.0) + p.amount
        payments_list.append({
            "id": p.id,
            "student_id": p.student_id,
            "student_name": s_name or "Noma'lum",
            "login_id": s_login or "N/A",
            "phone": s_phone or "",
            "amount": p.amount,
            "payment_method": method,
            "month_for": p.month_for,
            "created_at": p.created_at.strftime("%Y-%m-%d %H:%M") if p.created_at else "",
            "note": p.note or ""
        })

    # 2. Expenses in date range
    exp_stmt = (
        select(Expense)
        .where(
            Expense.date >= start_d,
            Expense.date <= end_d
        )
        .order_by(desc(Expense.date))
    )
    exp_res = await db.execute(exp_stmt)
    expenses = exp_res.scalars().all()

    total_expense = 0.0
    expenses_by_category: Dict[str, float] = {}
    expenses_list = []
    for exp in expenses:
        total_expense += exp.amount
        cat = exp.category or "Boshqa"
        expenses_by_category[cat] = expenses_by_category.get(cat, 0.0) + exp.amount
        expenses_list.append({
            "id": exp.id,
            "title": exp.title,
            "amount": exp.amount,
            "category": cat,
            "date": exp.date.strftime("%Y-%m-%d") if exp.date else "",
            "description": exp.description or ""
        })

    # 3. Teacher Payrolls in date range
    payroll_stmt = (
        select(TeacherPayroll, User.full_name)
        .outerjoin(User, TeacherPayroll.teacher_id == User.id)
        .where(
            func.date(TeacherPayroll.created_at) >= start_d,
            func.date(TeacherPayroll.created_at) <= end_d
        )
    )
    payroll_res = await db.execute(payroll_stmt)
    payroll_rows = payroll_res.all()
    total_teacher_salaries_paid = sum(r[0].paid_amount for r in payroll_rows)
    total_teacher_salaries_calculated = sum(r[0].calculated_salary for r in payroll_rows)

    net_profit = total_income - total_expense

    # 4. Students Statistics
    active_students_res = await db.execute(
        select(func.count(User.id)).where(User.role == UserRole.STUDENT, User.is_active == True)
    )
    total_active_students = active_students_res.scalar() or 0

    new_students_res = await db.execute(
        select(func.count(User.id)).where(
            User.role == UserRole.STUDENT,
            func.date(User.created_at) >= start_d,
            func.date(User.created_at) <= end_d
        )
    )
    new_students_count = new_students_res.scalar() or 0

    # 5. Lessons & Attendance in date range
    lessons_res = await db.execute(
        select(Lesson).where(Lesson.lesson_date >= start_d, Lesson.lesson_date <= end_d)
    )
    lessons_list = lessons_res.scalars().all()
    total_lessons_held = len(lessons_list)
    lesson_ids = [l.id for l in lessons_list]

    attendance_stats = {"PRESENT": 0, "LATE": 0, "ABSENT": 0, "EXCUSED": 0}
    if lesson_ids:
        att_res = await db.execute(
            select(Attendance.status, func.count(Attendance.id))
            .where(Attendance.lesson_id.in_(lesson_ids))
            .group_by(Attendance.status)
        )
        for st, count in att_res.all():
            st_val = st.value if hasattr(st, "value") else str(st)
            if st_val in attendance_stats:
                attendance_stats[st_val] = count

    total_att_records = sum(attendance_stats.values())
    present_total = attendance_stats["PRESENT"] + attendance_stats["LATE"]
    attendance_rate = round((present_total / total_att_records * 100), 1) if total_att_records > 0 else 100.0

    # 6. Exams & Results in date range
    exams_res = await db.execute(
        select(Exam).where(Exam.exam_date >= start_d, Exam.exam_date <= end_d)
    )
    exams = exams_res.scalars().all()
    exams_count = len(exams)
    exam_ids = [e.id for e in exams]

    exams_avg_score = 0.0
    exams_pass_rate = 0.0
    if exam_ids:
        res_stmt = select(func.avg(ExamResult.score), func.count(ExamResult.id)).where(ExamResult.exam_id.in_(exam_ids))
        res_calc = await db.execute(res_stmt)
        row_avg = res_calc.one_or_none()
        if row_avg and row_avg[0] is not None:
            exams_avg_score = round(float(row_avg[0]), 1)
            total_res_cnt = row_avg[1]
            if total_res_cnt > 0:
                passed_stmt = select(func.count(ExamResult.id)).where(
                    ExamResult.exam_id.in_(exam_ids),
                    ExamResult.score >= 70.0
                )
                passed_cnt = (await db.execute(passed_stmt)).scalar() or 0
                exams_pass_rate = round((passed_cnt / total_res_cnt) * 100, 1)

    # 7. Courses Performance Breakdown
    courses_res = await db.execute(select(Course).where(Course.is_active == True))
    all_courses = courses_res.scalars().all()
    courses_breakdown = []

    for c in all_courses:
        grp_res = await db.execute(select(Group).where(Group.course_id == c.id, Group.is_active == True))
        c_groups = grp_res.scalars().all()
        c_group_ids = [g.id for g in c_groups]

        students_cnt = 0
        if c_group_ids:
            st_count_res = await db.execute(
                select(func.count(GroupStudent.student_id))
                .where(GroupStudent.group_id.in_(c_group_ids), GroupStudent.is_active == True)
            )
            students_cnt = st_count_res.scalar() or 0

        courses_breakdown.append({
            "course_id": c.id,
            "title": c.title,
            "price_monthly": c.price_monthly,
            "duration_months": c.duration_months,
            "active_groups": len(c_groups),
            "total_students": students_cnt,
            "estimated_monthly_revenue": students_cnt * c.price_monthly
        })

    # 8. Teachers Performance Breakdown
    teachers_res = await db.execute(select(User).where(User.role == UserRole.TEACHER, User.is_active == True))
    teachers = teachers_res.scalars().all()
    teachers_breakdown = []

    for t in teachers:
        t_groups_res = await db.execute(select(Group).where(Group.teacher_id == t.id, Group.is_active == True))
        t_groups = t_groups_res.scalars().all()
        t_grp_ids = [g.id for g in t_groups]

        t_students_cnt = 0
        if t_grp_ids:
            st_cnt_res = await db.execute(
                select(func.count(GroupStudent.student_id))
                .where(GroupStudent.group_id.in_(t_grp_ids), GroupStudent.is_active == True)
            )
            t_students_cnt = st_cnt_res.scalar() or 0

        t_payroll_sum = sum(r[0].paid_amount for r in payroll_rows if r[0].teacher_id == t.id)

        teachers_breakdown.append({
            "teacher_id": t.id,
            "full_name": t.full_name,
            "phone": t.phone,
            "login_id": t.login_id,
            "groups_count": len(t_groups),
            "students_count": t_students_cnt,
            "salary_paid": t_payroll_sum
        })

    # 9. Time Trend (Monthly or Daily)
    trend_data = []
    if period_type == "yearly":
        year_val = start_d.year
        for m in range(1, 13):
            m_start = date(year_val, m, 1)
            m_last = calendar.monthrange(year_val, m)[1]
            m_end = date(year_val, m, m_last)

            m_pay = await db.execute(
                select(func.sum(Payment.amount)).where(
                    func.date(Payment.created_at) >= m_start,
                    func.date(Payment.created_at) <= m_end
                )
            )
            m_income = m_pay.scalar() or 0.0

            m_exp = await db.execute(
                select(func.sum(Expense.amount)).where(
                    Expense.date >= m_start,
                    Expense.date <= m_end
                )
            )
            m_expense = m_exp.scalar() or 0.0

            trend_data.append({
                "label": MONTH_NAMES_UZ.get(m, f"{m}-oy"),
                "month_num": m,
                "income": m_income,
                "expense": m_expense,
                "net_profit": m_income - m_expense
            })
    elif period_type == "monthly":
        num_days = (end_d - start_d).days + 1
        chunk_size = max(1, num_days // 4)
        cur = start_d
        while cur <= end_d:
            cur_end = min(end_d, cur + timedelta(days=chunk_size - 1))
            w_pay = await db.execute(
                select(func.sum(Payment.amount)).where(
                    func.date(Payment.created_at) >= cur,
                    func.date(Payment.created_at) <= cur_end
                )
            )
            w_exp = await db.execute(
                select(func.sum(Expense.amount)).where(
                    Expense.date >= cur,
                    Expense.date <= cur_end
                )
            )
            w_inc = w_pay.scalar() or 0.0
            w_ex = w_exp.scalar() or 0.0
            trend_data.append({
                "label": f"{cur.day}-{cur_end.day} {MONTH_NAMES_UZ.get(cur.month, '')}",
                "income": w_inc,
                "expense": w_ex,
                "net_profit": w_inc - w_ex
            })
            cur = cur_end + timedelta(days=1)
    else:  # Custom
        num_days = (end_d - start_d).days + 1
        steps = min(6, num_days)
        step_days = max(1, num_days // steps) if steps > 0 else 1
        cur = start_d
        while cur <= end_d:
            cur_end = min(end_d, cur + timedelta(days=step_days - 1))
            c_pay = await db.execute(
                select(func.sum(Payment.amount)).where(
                    func.date(Payment.created_at) >= cur,
                    func.date(Payment.created_at) <= cur_end
                )
            )
            c_exp = await db.execute(
                select(func.sum(Expense.amount)).where(
                    Expense.date >= cur,
                    Expense.date <= cur_end
                )
            )
            c_inc = c_pay.scalar() or 0.0
            c_ex = c_exp.scalar() or 0.0
            trend_data.append({
                "label": f"{cur.strftime('%d.%m')} - {cur_end.strftime('%d.%m')}",
                "income": c_inc,
                "expense": c_ex,
                "net_profit": c_inc - c_ex
            })
            cur = cur_end + timedelta(days=1)

    return {
        "period_type": period_type,
        "period_label": label,
        "start_date": start_d.strftime("%Y-%m-%d"),
        "end_date": end_d.strftime("%Y-%m-%d"),
        "financial": {
            "total_income": total_income,
            "total_expense": total_expense,
            "net_profit": net_profit,
            "payments_count": len(payments_list),
            "expenses_count": len(expenses_list),
            "by_payment_method": by_payment_method,
            "expenses_by_category": expenses_by_category,
            "teacher_salaries_paid": total_teacher_salaries_paid,
            "teacher_salaries_calculated": total_teacher_salaries_calculated
        },
        "academic_and_students": {
            "total_active_students": total_active_students,
            "new_students_count": new_students_count,
            "total_lessons_held": total_lessons_held,
            "attendance_rate": attendance_rate,
            "attendance_breakdown": attendance_stats,
            "exams_count": exams_count,
            "exams_avg_score": exams_avg_score,
            "exams_pass_rate": exams_pass_rate
        },
        "courses_breakdown": courses_breakdown,
        "teachers_breakdown": teachers_breakdown,
        "trend_data": trend_data,
        "payments": payments_list,
        "expenses": expenses_list
    }


# ----------------------------------------------------
# 1. SUMMARY JSON ENDPOINT
# ----------------------------------------------------
@router.get("/summary")
async def get_report_summary(
    period_type: str = Query("monthly", description="monthly | yearly | custom"),
    year: Optional[int] = Query(None),
    month: Optional[int] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin)
):
    start_d, end_d, start_dt, end_dt, label = resolve_date_range(period_type, year, month, start_date, end_date)
    data = await collect_report_data(db, start_d, end_d, start_dt, end_dt, period_type, label)
    return data


# ----------------------------------------------------
# 2. EXPORT TO EXCEL (.XLSX)
# ----------------------------------------------------
@router.get("/export/excel")
async def export_report_excel(
    period_type: str = Query("monthly"),
    year: Optional[int] = Query(None),
    month: Optional[int] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin)
):
    start_d, end_d, start_dt, end_dt, label = resolve_date_range(period_type, year, month, start_date, end_date)
    data = await collect_report_data(db, start_d, end_d, start_dt, end_dt, period_type, label)

    wb = openpyxl.Workbook()
    # Sheet 1: Xulosa
    ws_summary = wb.active
    ws_summary.title = "Umumiy Xulosa"

    # Styles
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="475569")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    kpi_title_font = Font(name="Calibri", size=10, bold=True, color="475569")
    kpi_val_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    kpi_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # Header section
    ws_summary["A1"] = "TA'LIM PLUS O'QUV MARKAZI — RASMIY HISOBOT"
    ws_summary["A1"].font = title_font
    ws_summary["A2"] = f"Hisobot Davri: {data['period_label']} ({data['start_date']} — {data['end_date']})"
    ws_summary["A2"].font = subtitle_font
    ws_summary["A3"] = f"Generatsiya vaqti: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Mas'ul: {admin.full_name}"
    ws_summary["A3"].font = subtitle_font

    # KPI Boxes
    kpi_data = [
        ("Jami Tushum (Daromad)", f"{data['financial']['total_income']:,.0f} UZS"),
        ("Jami Xarajatlar", f"{data['financial']['total_expense']:,.0f} UZS"),
        ("Sof Foyda", f"{data['financial']['net_profit']:,.0f} UZS"),
        ("Faol O'quvchilar", f"{data['academic_and_students']['total_active_students']} ta"),
        ("Yangi O'quvchilar", f"{data['academic_and_students']['new_students_count']} ta"),
        ("O'rtacha Davomat", f"{data['academic_and_students']['attendance_rate']}%")
    ]

    r = 5
    for idx, (k_title, k_val) in enumerate(kpi_data):
        col_start = (idx % 3) * 2 + 1
        row_pos = r if idx < 3 else r + 3
        
        c_cell = ws_summary.cell(row=row_pos, column=col_start, value=k_title)
        c_cell.font = kpi_title_font
        c_cell.fill = kpi_fill
        c_cell.border = border_thin
        
        v_cell = ws_summary.cell(row=row_pos + 1, column=col_start, value=k_val)
        v_cell.font = kpi_val_font
        v_cell.fill = kpi_fill
        v_cell.border = border_thin

    # Kurslar jadvali
    ws_summary.cell(row=12, column=1, value="Kurslar Kesimida Natijalar").font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    headers_courses = ["Kurs Nomi", "Oylik Narxi (UZS)", "Davomiyligi", "Guruhlar Soni", "O'quvchilar Soni", "Oylik Prognoz Tushum"]
    for c_idx, h in enumerate(headers_courses, start=1):
        cell = ws_summary.cell(row=13, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin

    row_c = 14
    for crs in data["courses_breakdown"]:
        ws_summary.cell(row=row_c, column=1, value=crs["title"]).border = border_thin
        c2 = ws_summary.cell(row=row_c, column=2, value=crs["price_monthly"])
        c2.number_format = "#,##0"
        c2.border = border_thin
        ws_summary.cell(row=row_c, column=3, value=f"{crs['duration_months']} oy").border = border_thin
        ws_summary.cell(row=row_c, column=4, value=crs["active_groups"]).border = border_thin
        ws_summary.cell(row=row_c, column=5, value=crs["total_students"]).border = border_thin
        c6 = ws_summary.cell(row=row_c, column=6, value=crs["estimated_monthly_revenue"])
        c6.number_format = "#,##0"
        c6.border = border_thin
        row_c += 1

    # To'lov turlari va Xarajat toifalari
    row_sub = row_c + 2
    ws_summary.cell(row=row_sub, column=1, value="To'lov Usullari Taqsimoti").font = Font(name="Calibri", size=12, bold=True)
    ws_summary.cell(row=row_sub, column=4, value="Xarajatlar Toifalari").font = Font(name="Calibri", size=12, bold=True)
    row_sub += 1

    ws_summary.cell(row=row_sub, column=1, value="To'lov Turi").font = header_font
    ws_summary.cell(row=row_sub, column=1).fill = header_fill
    ws_summary.cell(row=row_sub, column=2, value="Summa (UZS)").font = header_font
    ws_summary.cell(row=row_sub, column=2).fill = header_fill

    ws_summary.cell(row=row_sub, column=4, value="Kategoriya").font = header_font
    ws_summary.cell(row=row_sub, column=4).fill = header_fill
    ws_summary.cell(row=row_sub, column=5, value="Summa (UZS)").font = header_font
    ws_summary.cell(row=row_sub, column=5).fill = header_fill

    r_pay = row_sub + 1
    for k, v in data["financial"]["by_payment_method"].items():
        ws_summary.cell(row=r_pay, column=1, value=k).border = border_thin
        cp = ws_summary.cell(row=r_pay, column=2, value=v)
        cp.number_format = "#,##0"
        cp.border = border_thin
        r_pay += 1

    r_exp = row_sub + 1
    for k, v in data["financial"]["expenses_by_category"].items():
        ws_summary.cell(row=r_exp, column=4, value=k).border = border_thin
        ce = ws_summary.cell(row=r_exp, column=5, value=v)
        ce.number_format = "#,##0"
        ce.border = border_thin
        r_exp += 1

    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # Sheet 2: To'lovlar Ro'yxati
    ws_payments = wb.create_sheet(title="To'lovlar Ro'yxati")
    ws_payments["A1"] = f"Barcha Qabul Qilingan To'lovlar ({data['period_label']})"
    ws_payments["A1"].font = title_font

    pay_headers = ["№", "Chek ID", "Sana", "O'quvchi Ismi", "Login ID", "Telefon", "Summa (UZS)", "To'lov Turi", "Oy Uchun", "Izoh"]
    for c_idx, h in enumerate(pay_headers, start=1):
        cell = ws_payments.cell(row=3, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin

    for idx, p in enumerate(data["payments"], start=1):
        r_num = idx + 3
        ws_payments.cell(row=r_num, column=1, value=idx).border = border_thin
        ws_payments.cell(row=r_num, column=2, value=f"#{p['id']}").border = border_thin
        ws_payments.cell(row=r_num, column=3, value=p["created_at"]).border = border_thin
        ws_payments.cell(row=r_num, column=4, value=p["student_name"]).border = border_thin
        ws_payments.cell(row=r_num, column=5, value=p["login_id"]).border = border_thin
        ws_payments.cell(row=r_num, column=6, value=p["phone"]).border = border_thin
        amt_cell = ws_payments.cell(row=r_num, column=7, value=p["amount"])
        amt_cell.number_format = "#,##0"
        amt_cell.border = border_thin
        ws_payments.cell(row=r_num, column=8, value=p["payment_method"]).border = border_thin
        ws_payments.cell(row=r_num, column=9, value=p["month_for"]).border = border_thin
        ws_payments.cell(row=r_num, column=10, value=p["note"]).border = border_thin

    for col in ws_payments.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_payments.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Sheet 3: Xarajatlar
    ws_expenses = wb.create_sheet(title="Xarajatlar Ro'yxati")
    ws_expenses["A1"] = f"Markaz Xarajatlari ({data['period_label']})"
    ws_expenses["A1"].font = title_font

    exp_headers = ["№", "ID", "Sana", "Xarajat Nomi", "Kategoriya", "Summa (UZS)", "Tavsif"]
    for c_idx, h in enumerate(exp_headers, start=1):
        cell = ws_expenses.cell(row=3, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin

    for idx, e in enumerate(data["expenses"], start=1):
        r_num = idx + 3
        ws_expenses.cell(row=r_num, column=1, value=idx).border = border_thin
        ws_expenses.cell(row=r_num, column=2, value=e["id"]).border = border_thin
        ws_expenses.cell(row=r_num, column=3, value=e["date"]).border = border_thin
        ws_expenses.cell(row=r_num, column=4, value=e["title"]).border = border_thin
        ws_expenses.cell(row=r_num, column=5, value=e["category"]).border = border_thin
        e_amt = ws_expenses.cell(row=r_num, column=6, value=e["amount"])
        e_amt.number_format = "#,##0"
        e_amt.border = border_thin
        ws_expenses.cell(row=r_num, column=7, value=e["description"]).border = border_thin

    for col in ws_expenses.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_expenses.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Sheet 4: O'qituvchilar
    ws_teachers = wb.create_sheet(title="O'qituvchilar Faoliyati")
    ws_teachers["A1"] = "O'qituvchilar va Guruhlar Samaradorligi"
    ws_teachers["A1"].font = title_font

    tch_headers = ["№", "O'qituvchi F.I.Sh", "Login ID", "Telefon", "Faol Guruhlar", "O'quvchilar Soni", "To'langan Oylik (UZS)"]
    for c_idx, h in enumerate(tch_headers, start=1):
        cell = ws_teachers.cell(row=3, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = border_thin

    for idx, t in enumerate(data["teachers_breakdown"], start=1):
        r_num = idx + 3
        ws_teachers.cell(row=r_num, column=1, value=idx).border = border_thin
        ws_teachers.cell(row=r_num, column=2, value=t["full_name"]).border = border_thin
        ws_teachers.cell(row=r_num, column=3, value=t["login_id"]).border = border_thin
        ws_teachers.cell(row=r_num, column=4, value=t["phone"]).border = border_thin
        ws_teachers.cell(row=r_num, column=5, value=t["groups_count"]).border = border_thin
        ws_teachers.cell(row=r_num, column=6, value=t["students_count"]).border = border_thin
        t_sal = ws_teachers.cell(row=r_num, column=7, value=t["salary_paid"])
        t_sal.number_format = "#,##0"
        t_sal.border = border_thin

    for col in ws_teachers.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_teachers.column_dimensions[col_letter].width = max(max_len + 3, 14)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"hisobot_{start_d}_{end_d}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ----------------------------------------------------
# 3. EXPORT TO PDF (.PDF)
# ----------------------------------------------------
@router.get("/export/pdf")
async def export_report_pdf(
    period_type: str = Query("monthly"),
    year: Optional[int] = Query(None),
    month: Optional[int] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin)
):
    start_d, end_d, start_dt, end_dt, label = resolve_date_range(period_type, year, month, start_date, end_date)
    data = await collect_report_data(db, start_d, end_d, start_dt, end_dt, period_type, label)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=1
    )
    sub_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#475569'),
        alignment=1
    )
    section_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#1E40AF'),
        spaceBefore=12,
        spaceAfter=6
    )
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=11,
        textColor=colors.HexColor('#0F172A')
    )
    cell_bold = ParagraphStyle(
        'CellBold',
        parent=cell_style,
        fontName='Helvetica-Bold'
    )
    cell_right = ParagraphStyle(
        'CellRight',
        parent=cell_style,
        alignment=2
    )

    story = []

    story.append(Paragraph("TA'LIM PLUS O'QUV MARKAZI", title_style))
    story.append(Paragraph("DAVRIY RASMIY FAOLIYAT VA MOLIYA HISOBOTI", ParagraphStyle('SubSub', parent=sub_style, fontName='Helvetica-Bold', fontSize=12, textColor=colors.HexColor('#2563EB'))))
    story.append(Paragraph(f"Davr: <b>{data['period_label']}</b> ({data['start_date']} — {data['end_date']})", sub_style))
    story.append(Paragraph(f"Chop etilgan sana: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Mas'ul shaxs: {admin.full_name}", sub_style))
    story.append(Spacer(1, 12))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2563EB'), spaceAfter=14))

    # 1. KPI
    story.append(Paragraph("1. Asosiy Ko'rsatkichlar (KPI)", section_style))
    kpi_table_data = [
        [
            Paragraph("<b>Jami Tushum (Daromad):</b>", cell_style),
            Paragraph(f"<b>{data['financial']['total_income']:,.0f} UZS</b>", cell_right),
            Paragraph("<b>Faol O'quvchilar:</b>", cell_style),
            Paragraph(f"<b>{data['academic_and_students']['total_active_students']} nafar</b>", cell_right)
        ],
        [
            Paragraph("<b>Jami Xarajatlar:</b>", cell_style),
            Paragraph(f"<b>{data['financial']['total_expense']:,.0f} UZS</b>", cell_right),
            Paragraph("<b>Yangi O'quvchilar:</b>", cell_style),
            Paragraph(f"<b>+{data['academic_and_students']['new_students_count']} nafar</b>", cell_right)
        ],
        [
            Paragraph("<b>Sof Foyda:</b>", cell_style),
            Paragraph(f"<font color='{'#16A34A' if data['financial']['net_profit'] >= 0 else '#DC2626'}'><b>{data['financial']['net_profit']:,.0f} UZS</b></font>", cell_right),
            Paragraph("<b>O'rtacha Davomat:</b>", cell_style),
            Paragraph(f"<b>{data['academic_and_students']['attendance_rate']}%</b>", cell_right)
        ],
        [
            Paragraph("<b>O'qituvchilarga To'langan Oylik:</b>", cell_style),
            Paragraph(f"<b>{data['financial']['teacher_salaries_paid']:,.0f} UZS</b>", cell_right),
            Paragraph("<b>O'tkazilgan Darslar:</b>", cell_style),
            Paragraph(f"<b>{data['academic_and_students']['total_lessons_held']} ta</b>", cell_right)
        ]
    ]

    t_kpi = Table(kpi_table_data, colWidths=[150, 110, 150, 110])
    t_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(t_kpi)
    story.append(Spacer(1, 14))

    # 2. Kurslar
    story.append(Paragraph("2. Kurslar Faoliyati va Tushumlar", section_style))
    course_table_data = [
        [
            Paragraph("<b>Kurs Nomi</b>", cell_bold),
            Paragraph("<b>Oylik Narxi</b>", cell_bold),
            Paragraph("<b>Guruhlar</b>", cell_bold),
            Paragraph("<b>O'quvchilar</b>", cell_bold),
            Paragraph("<b>Oylik Prognoz</b>", cell_bold)
        ]
    ]
    for crs in data["courses_breakdown"]:
        course_table_data.append([
            Paragraph(crs["title"], cell_style),
            Paragraph(f"{crs['price_monthly']:,.0f} UZS", cell_style),
            Paragraph(str(crs["active_groups"]), cell_style),
            Paragraph(str(crs["total_students"]), cell_style),
            Paragraph(f"{crs['estimated_monthly_revenue']:,.0f} UZS", cell_right)
        ])

    t_courses = Table(course_table_data, colWidths=[160, 90, 60, 70, 140])
    t_courses.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E40AF')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t_courses)
    story.append(Spacer(1, 14))

    # 3. To'lov Usullari va Xarajatlar
    story.append(Paragraph("3. To'lov Usullari va Xarajatlar Kategoriyalari", section_style))
    pay_methods_data = [[Paragraph("<b>To'lov Turi</b>", cell_bold), Paragraph("<b>Summa (UZS)</b>", cell_bold)]]
    for k, v in data["financial"]["by_payment_method"].items():
        pay_methods_data.append([Paragraph(k, cell_style), Paragraph(f"{v:,.0f} UZS", cell_right)])

    exp_cats_data = [[Paragraph("<b>Xarajat Toifasi</b>", cell_bold), Paragraph("<b>Summa (UZS)</b>", cell_bold)]]
    for k, v in data["financial"]["expenses_by_category"].items():
        exp_cats_data.append([Paragraph(k, cell_style), Paragraph(f"{v:,.0f} UZS", cell_right)])
    if not data["financial"]["expenses_by_category"]:
        exp_cats_data.append([Paragraph("Xarajat kiritilmagan", cell_style), Paragraph("0 UZS", cell_right)])

    t_pay = Table(pay_methods_data, colWidths=[140, 110])
    t_pay.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))

    t_exp = Table(exp_cats_data, colWidths=[140, 110])
    t_exp.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DC2626')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))

    side_by_side = Table([[t_pay, Paragraph("", cell_style), t_exp]], colWidths=[250, 20, 250])
    side_by_side.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
    ]))
    story.append(side_by_side)
    story.append(Spacer(1, 24))

    # 4. Imzo va Muhr
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#94A3B8'), spaceAfter=14))
    sign_data = [
        [
            Paragraph("<b>O'quv Markazi Rahbari:</b>", cell_bold),
            Paragraph("_________________ (imzo)", cell_style),
            Paragraph("<b>Bosh Hisobchi:</b>", cell_bold),
            Paragraph("_________________ (imzo)", cell_style)
        ],
        [
            Paragraph(f"F.I.Sh: <b>{admin.full_name}</b>", cell_style),
            Paragraph("M.O' (Muhr o'rni)", cell_style),
            Paragraph("F.I.Sh: _________________", cell_style),
            Paragraph("Sana: " + datetime.now().strftime("%d.%m.%Y"), cell_style)
        ]
    ]
    t_sign = Table(sign_data, colWidths=[150, 110, 140, 120])
    t_sign.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t_sign)

    doc.build(story)
    buffer.seek(0)

    filename = f"hisobot_{start_d}_{end_d}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ----------------------------------------------------
# 4. EXPORT TO WORD (.DOCX)
# ----------------------------------------------------
@router.get("/export/docx")
async def export_report_docx(
    period_type: str = Query("monthly"),
    year: Optional[int] = Query(None),
    month: Optional[int] = Query(None),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: AsyncSession = Depends(get_db),
    admin: User = Depends(get_current_admin)
):
    start_d, end_d, start_dt, end_dt, label = resolve_date_range(period_type, year, month, start_date, end_date)
    data = await collect_report_data(db, start_d, end_d, start_dt, end_dt, period_type, label)

    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("TA'LIM PLUS O'QUV MARKAZI\nRASMIY FAOLIYAT HISOBOTI")
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run(f"Davr: {data['period_label']} ({data['start_date']} — {data['end_date']})\nTayyorlagan: {admin.full_name} | Vaqt: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 1. KPI
    h1 = doc.add_heading("1. Asosiy Ko'rsatkichlar (KPI)", level=2)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    kpi_table = doc.add_table(rows=4, cols=2)
    kpi_table.style = 'Table Grid'
    kpi_rows = [
        ("Jami Tushum (Daromad):", f"{data['financial']['total_income']:,.0f} UZS"),
        ("Jami Xarajatlar:", f"{data['financial']['total_expense']:,.0f} UZS"),
        ("Sof Foyda:", f"{data['financial']['net_profit']:,.0f} UZS"),
        ("Faol O'quvchilar Soni:", f"{data['academic_and_students']['total_active_students']} nafar (Yangi: +{data['academic_and_students']['new_students_count']})")
    ]
    for idx, (label_txt, val_txt) in enumerate(kpi_rows):
        row_cells = kpi_table.rows[idx].cells
        row_cells[0].text = label_txt
        row_cells[0].paragraphs[0].runs[0].font.bold = True
        row_cells[1].text = val_txt

    # 2. Kurslar
    h2 = doc.add_heading("2. Kurslar Faoliyati", level=2)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    course_table = doc.add_table(rows=1 + len(data["courses_breakdown"]), cols=5)
    course_table.style = 'Table Grid'
    hdr_cells = course_table.rows[0].cells
    hdr_titles = ["Kurs Nomi", "Oylik Narxi", "Davomiyligi", "Guruhlar", "O'quvchilar"]
    for i, title in enumerate(hdr_titles):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True

    for i, crs in enumerate(data["courses_breakdown"], start=1):
        row_cells = course_table.rows[i].cells
        row_cells[0].text = crs["title"]
        row_cells[1].text = f"{crs['price_monthly']:,.0f} UZS"
        row_cells[2].text = f"{crs['duration_months']} oy"
        row_cells[3].text = str(crs["active_groups"])
        row_cells[4].text = str(crs["total_students"])

    # 3. Signatures
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    p_sign = doc.add_paragraph()
    p_sign.add_run("Rahbar imzosi: _________________________\t\t\tMuhr o'rni: (M.O')\n\nBosh hisobchi: _________________________")

    output = BytesIO()
    doc.save(output)
    output.seek(0)

    filename = f"hisobot_{start_d}_{end_d}.docx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
